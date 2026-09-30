"""Tests for Unnamed Defense (Resposta à Acusação Inominada) and Unified Witness Deduplication & Labeling.

Desenvolvido por FChNeto.
"""

import io
from pathlib import Path
import pytest
import pymupdf
import docx

from app.core.models import (
    Defendant,
    HearingSummaryData,
    HistoryItem,
    PJeDocument,
    Witness,
)
from app.engines.offline_engine import OfflineExtractionEngine
from app.generators.docx_generator import generate_docx_summary


def test_unnamed_defense_detection_and_synthesis():
    """Validates detection of an unnamed petition (Petição / Manifestação) containing Resposta à Acusação."""
    engine = OfflineExtractionEngine()

    p_doc = PJeDocument(
        doc_id="9988776",
        date_str="14/03/2025 10:30",
        doc_name="Petição",
        doc_type="Petição (Outras)",
        start_page=1,
        end_page=2,
    )

    doc_text = """
    EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA VARA CRIMINAL
    Autos nº 0801234-56.2025.8.20.5001

    MARCOS VINICIUS DA SILVA, já qualificado nos autos da AÇÃO PENAL que lhe move o MINISTÉRIO PÚBLICO,
    vem, por intermédio de seu advogado que esta subscreve, com fundamento nos artigos 396 e 396-A do CPP,
    apresentar sua DEFESA PRELIMINAR / MANIFESTAÇÃO.

    1. PRELIMINARMENTE - DA INÉPCIA DA DENÚNCIA
    A denúncia é manifestamente inepta por não individualizar as condutas.

    2. DO MÉRITO - DA ABSOLVIÇÃO SUMÁRIA
    Requer a absolvição sumária nos termos do art. 397, III, do CPP, por ausência de dolo.

    3. DO ROL DE TESTEMUNHAS
    Requer a oitiva das mesmas testemunhas arroladas na denúncia pelo Ministério Público.

    Nestes termos, pede deferimento.
    Dr. Roberto Alencar - OAB/RN 15420
    Advogado
    """

    reus_capa = ["MARCOS VINICIUS DA SILVA"]
    advs_capa = ["ROBERTO ALENCAR"]

    parsed = engine._parse_defense_document(p_doc, doc_text, reus_capa, advs_capa, is_named_defense=False)
    assert parsed is not None, "Failed to identify unnamed petition as defense response"
    assert parsed["is_inominada"] is True
    assert parsed["is_defensoria"] is False
    assert "Dr. Roberto Alencar" in parsed["petitioner_str"]
    assert "15420" in parsed["petitioner_str"]
    assert "Marcos Vinicius Da Silva" in parsed["represented_defendants"][0]
    assert parsed["adopts_mp_witnesses"] is True
    assert "inépcia" in parsed["summary_for_history"].lower()
    assert "absolvição sumária" in parsed["summary_for_history"].lower()


def test_multi_defendants_with_distinct_defenses():
    """Validates multi-defendant case with distinct defense responses (Defensoria and Private Lawyer)."""
    engine = OfflineExtractionEngine()

    reus_capa = ["BRUNO CESAR LIMA", "DIEGO TAVARES GOMES"]
    advs_capa = ["DEFENSORIA PUBLICA", "LUCAS MEDEIROS"]

    p_doc1 = PJeDocument(
        doc_id="1001",
        date_str="10/02/2025",
        doc_name="Manifestação",
        doc_type="Petição",
        start_page=1,
        end_page=1,
    )
    txt1 = """
    A DEFENSORIA PÚBLICA DO ESTADO DO RIO GRANDE DO NORTE, em defesa do assistido BRUNO CESAR LIMA,
    vem apresentar RESPOSTA À ACUSAÇÃO (art. 396 do CPP), pugnando pela absolvição sumária.
    Reitera o rol de testemunhas da denúncia.
    """

    p_doc2 = PJeDocument(
        doc_id="1002",
        date_str="12/02/2025",
        doc_name="Petição Avulsa",
        doc_type="Petição",
        start_page=2,
        end_page=2,
    )
    txt2 = """
    DIEGO TAVARES GOMES, por seu advogado Dr. Lucas Medeiros - OAB/RN 9876, apresenta resposta à acusação
    (art. 396-A CPP), reservando-se para o mérito em alegações finais.
    ROL DE TESTEMUNHAS:
    1. Francisco das Chagas
    Dr. Lucas Medeiros - OAB/RN 9876
    """

    parsed1 = engine._parse_defense_document(p_doc1, txt1, reus_capa, advs_capa, False)
    parsed2 = engine._parse_defense_document(p_doc2, txt2, reus_capa, advs_capa, False)

    assert parsed1 is not None and parsed2 is not None
    assert parsed1["is_defensoria"] is True
    assert "Bruno Cesar Lima" in parsed1["represented_defendants"][0]
    assert parsed2["is_defensoria"] is False
    assert "Diego Tavares Gomes" in parsed2["represented_defendants"][0]
    assert "Lucas Medeiros" in parsed2["lawyer_name"]

    defense_docs = [parsed1, parsed2]
    counsel_str = engine._extract_defense_counsel("", advs_capa, None, defense_docs, reus_capa)
    assert "Assistido pela Defensoria Pública" in counsel_str
    assert "Bruno Cesar Lima" in counsel_str
    assert "Dr. Lucas Medeiros" in counsel_str
    assert "Diego Tavares Gomes" in counsel_str


def test_witness_deduplication_and_unified_labeling_defensoria():
    """Validates unified witness labeling when defense is Defensoria Pública."""
    engine = OfflineExtractionEngine()

    denuncia_text = """
    ROL DE TESTEMUNHAS:
    1. Tenente Silva (Policial Militar)
    2. Maria Joana (vítima)
    3. Carlos Alberto (Testemunha)
    """

    defense_doc_info = {
        "doc_id": "5555",
        "is_defensoria": True,
        "lawyer_name": "",
        "represented_defendants": ["João da Silva"],
        "adopts_mp_witnesses": True,
        "explicit_witnesses": [],
    }

    pros_w, def_w, def_note = engine._extract_witnesses(
        denuncia_text=denuncia_text,
        resposta_text="",
        vitimas_capa=[],
        testemunhas_capa=[],
        mandados_docs=[],
        doc=None,
        defense_docs=[defense_doc_info],
    )

    assert len(pros_w) == 3
    assert len(def_w) == 0
    for w in pros_w:
        assert "(arrolada pelo Ministério Público e pela Defensoria Pública)" in w.role


def test_witness_deduplication_and_unified_labeling_private_lawyer():
    """Validates unified witness labeling when defense is private lawyer."""
    engine = OfflineExtractionEngine()

    denuncia_text = """
    ROL DE TESTEMUNHAS:
    1. Soldado Souza (Policial Militar)
    2. Ana Paula (vítima)
    """

    defense_doc_info = {
        "doc_id": "7777",
        "is_defensoria": False,
        "lawyer_name": "Renato Garcia",
        "lawyer_oab": "OAB/RN 11223",
        "represented_defendants": ["Pedro Rocha"],
        "adopts_mp_witnesses": True,
        "explicit_witnesses": [],
    }

    pros_w, def_w, def_note = engine._extract_witnesses(
        denuncia_text=denuncia_text,
        resposta_text="",
        vitimas_capa=[],
        testemunhas_capa=[],
        mandados_docs=[],
        doc=None,
        defense_docs=[defense_doc_info],
    )

    assert len(pros_w) == 2
    for w in pros_w:
        assert "(arrolada pelo Ministério Público e pelo Advogado Dr. Renato Garcia, em defesa de Pedro Rocha)" in w.role


def test_witness_deduplication_and_unified_labeling_all_parties():
    """Validates '(arrolada por todos)' when prosecution and all multiple co-defenses roll the witness."""
    engine = OfflineExtractionEngine()

    denuncia_text = """
    ROL DE TESTEMUNHAS:
    1. Sargento Pereira (Policial Militar)
    """

    defense_docs = [
        {
            "doc_id": "111",
            "is_defensoria": True,
            "represented_defendants": ["Réu 1"],
            "adopts_mp_witnesses": True,
            "explicit_witnesses": [],
        },
        {
            "doc_id": "222",
            "is_defensoria": False,
            "lawyer_name": "Advogado X",
            "represented_defendants": ["Réu 2"],
            "adopts_mp_witnesses": True,
            "explicit_witnesses": [],
        },
    ]

    pros_w, def_w, def_note = engine._extract_witnesses(
        denuncia_text=denuncia_text,
        resposta_text="",
        vitimas_capa=[],
        testemunhas_capa=[],
        mandados_docs=[],
        doc=None,
        defense_docs=defense_docs,
    )

    assert len(pros_w) == 1
    assert "(arrolada por todos)" in pros_w[0].role


def test_docx_generation_with_unified_witnesses_and_unnamed_defense():
    """Validates complete OpenXML DOCX document generation with unnamed defense and deduplicated witnesses."""
    data = HearingSummaryData(
        case_number="0809999-11.2025.8.20.5001",
        act_type="AIJ",
        hearing_datetime="15.10.26 às 14h00min",
        hearing_link="https://teams.microsoft.com/meet/999888777",
        is_in_person=False,
        prosecutor="Dr. Promotor de Justiça",
        defendants=[
            Defendant(name="MARCOS VINICIUS", status="respondendo em liberdade", subpoena_id="1234567"),
            Defendant(name="DIEGO TAVARES", status="réu preso", subpoena_id="1234568"),
        ],
        defense_counsel="Réu Marcos Vinicius assistido pela Defensoria Pública; Réu Diego Tavares representado por advogado particular, Dr. Lucas Medeiros - OAB/RN 9876",
        qualification_text="MARCOS VINICIUS, brasileiro, solteiro.\n\nDIEGO TAVARES, brasileiro, solteiro.",
        imputation_text="art. 157, § 2º, II, do CP",
        facts_summary="Consta na denúncia que os acusados subtraíram bens da vítima.",
        chronological_history=[
            HistoryItem(
                date_str="01/02/25",
                description="Denúncia oferecida pelo Ministério Público",
                doc_id="10001",
            ),
            HistoryItem(
                date_str="10/02/25",
                description="Decisão recebendo a denúncia",
                doc_id="10002",
            ),
            HistoryItem(
                date_str="20/02/25",
                description="Resposta à acusação apresentada pela Defensoria Pública em defesa de Marcos Vinicius, pugnando pela absolvição sumária e requerendo a oitiva das mesmas testemunhas da acusação",
                doc_id="10003",
            ),
            HistoryItem(
                date_str="25/02/25",
                description="Resposta à acusação apresentada pelo Advogado Dr. Lucas Medeiros (OAB/RN 9876) em defesa de Diego Tavares, arguindo preliminar de inépcia da denúncia e arrolando testemunhas",
                doc_id="10004",
            ),
        ],
        prosecution_witnesses=[
            Witness(
                number=1,
                name="Sargento Ribeiro",
                role="PM (arrolada por todos)",
                status_id="Ofício enviado ID 555666",
            ),
            Witness(
                number=2,
                name="Vitima Carla",
                role="vítima (arrolada por todos)",
                status_id="Intimada ID 555667",
            ),
        ],
        defense_witnesses=[
            Witness(
                number=1,
                name="Testemunha Exclusiva Defesa",
                role="Testemunha de Defesa",
                status_id="Intimado ID 555668",
            )
        ],
        defense_witness_note=None,
        closure_text="Cordial e respeitosamente,",
    )

    docx_bytes = generate_docx_summary(data)
    assert len(docx_bytes) > 1000

    # Parse generated docx
    doc_in = docx.Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(p.text for p in doc_in.paragraphs)

    assert "0809999-11.2025.8.20.5001" in all_text
    assert "MARCOS VINICIUS" in all_text
    assert "DIEGO TAVARES" in all_text
    assert "Resposta à acusação apresentada pela Defensoria Pública" in all_text
    assert "Resposta à acusação apresentada pelo Advogado Dr. Lucas Medeiros" in all_text
    assert "Sargento Ribeiro - PM (arrolada por todos)" in all_text
    assert "Testemunha Exclusiva Defesa" in all_text
    assert "Cordial e respeitosamente," in all_text


def test_multi_defendant_partial_witness_merge_combination():
    """Validates that when 2 out of 3 defenses adopt MP witnesses, both are included in the witness tag."""
    engine = OfflineExtractionEngine()

    denuncia_text = """
    ROL DE TESTEMUNHAS:
    1. Sargento Pereira (Policial Militar)
    2. Maria Joana (Vítima)
    """

    defense_docs = [
        {
            "doc_id": "111",
            "is_defensoria": True,
            "represented_defendants": ["Réu 1"],
            "adopts_mp_witnesses": True,
            "explicit_witnesses": [],
        },
        {
            "doc_id": "222",
            "is_defensoria": False,
            "lawyer_name": "Lucas Medeiros",
            "represented_defendants": ["Réu 2"],
            "adopts_mp_witnesses": True,
            "explicit_witnesses": [],
        },
        {
            "doc_id": "333",
            "is_defensoria": False,
            "lawyer_name": "Maria Silva",
            "represented_defendants": ["Réu 3"],
            "adopts_mp_witnesses": False,
            "explicit_witnesses": ["Testemunha Exclusiva 3"],
        },
    ]

    pros_w, def_w, def_note = engine._extract_witnesses(
        denuncia_text=denuncia_text,
        resposta_text="",
        vitimas_capa=[],
        testemunhas_capa=[],
        mandados_docs=[],
        doc=None,
        defense_docs=defense_docs,
    )

    assert len(pros_w) == 2
    assert len(def_w) == 1
    # Check that both Defensoria (Réu 1) and Dr. Lucas Medeiros (Réu 2) are in the tag, not dropped:
    for w in pros_w:
        assert "pela Defensoria Pública, em defesa de Réu 1" in w.role
        assert "Dr. Lucas Medeiros, em defesa de Réu 2" in w.role

    # Check exclusive defense witness
    assert def_w[0].name == "Testemunha Exclusiva 3"


def test_unnamed_defense_with_contestacao_and_documento_diverso():
    """Validates detection when document is labeled 'Documento Diverso' or 'Contestação'."""
    engine = OfflineExtractionEngine()

    p_doc = PJeDocument(
        doc_id="888111",
        date_str="05/04/2025 14:00",
        doc_name="Documento Diverso",
        doc_type="Documento Comprobatório",
        start_page=1,
        end_page=2,
    )

    txt = """
    EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO
    GABRIEL SANTOS, já qualificado, por seu advogado infra-assinado, vem apresentar RESPOSTA À ACUSAÇÃO,
    com fulcro no artigo 396-A do Código de Processo Penal.
    Requer a absolvição sumária por atipicidade da conduta (art. 397, III, CPP).
    Dr. Carlos Eduardo - OAB/RN 7788
    """

    reus_capa = ["GABRIEL SANTOS"]
    advs_capa = ["CARLOS EDUARDO"]

    parsed = engine._parse_defense_document(p_doc, txt, reus_capa, advs_capa, is_named_defense=False)
    assert parsed is not None
    assert "Carlos Eduardo" in parsed["lawyer_name"]
    assert "7788" in parsed["petitioner_str"]
    assert "Gabriel Santos" in parsed["represented_defendants"][0]
    assert "absolvição sumária" in parsed["summary_for_history"].lower()


def test_defense_lawyer_name_dr_prefix_no_duplicate():
    """Validates that 'Dr.' prefix in lawyer name is not duplicated into 'Dr. Dr.'."""
    engine = OfflineExtractionEngine()

    p_doc = PJeDocument(
        doc_id="999000",
        date_str="01/05/2025 10:00",
        doc_name="Petição",
        doc_type="Petição",
        start_page=1,
        end_page=1,
    )
    txt = """
    RESPOSTA À ACUSAÇÃO (art. 396 do CPP).
    Pugna pela absolvição sumária e reitera o rol de testemunhas da acusação.
    DR. FRANCISCO CHAGAS - OAB/RN 1234
    """
    reus_capa = ["JOÃO DA SILVA"]
    advs_capa = ["DR. FRANCISCO CHAGAS"]

    parsed = engine._parse_defense_document(p_doc, txt, reus_capa, advs_capa, is_named_defense=True)
    assert parsed is not None
    assert "Dr. Dr." not in parsed["petitioner_str"]
    assert "Dr. Francisco Chagas" in parsed["petitioner_str"]

    pros_w, def_w, _ = engine._extract_witnesses(
        denuncia_text="ROL DE TESTEMUNHAS:\n1. Soldado Silva (PM)",
        resposta_text="",
        vitimas_capa=[],
        testemunhas_capa=[],
        mandados_docs=[],
        doc=None,
        defense_docs=[parsed],
    )
    assert "Dr. Dr." not in pros_w[0].role
    assert "Dr. Francisco Chagas" in pros_w[0].role
    assert "em defesa de do réu" not in pros_w[0].role


def test_multiple_private_lawyers_counsel_string():
    """Validates defense counsel formatting when multiple co-defendants have distinct private attorneys."""
    engine = OfflineExtractionEngine()
    reus_capa = ["RÉU ALFA", "RÉU BETA"]
    advs_capa = ["DR. ADVOGADO UM", "DR. ADVOGADO DOIS"]

    defense_docs = [
        {
            "doc_id": "1",
            "is_defensoria": False,
            "lawyer_name": "Advogado Um",
            "lawyer_oab": "OAB/RN 111",
            "represented_defendants": ["Réu Alfa"],
        },
        {
            "doc_id": "2",
            "is_defensoria": False,
            "lawyer_name": "Advogado Dois",
            "lawyer_oab": "OAB/RN 222",
            "represented_defendants": ["Réu Beta"],
        },
    ]

    counsel = engine._extract_defense_counsel("", advs_capa, None, defense_docs, reus_capa)
    assert "Representados por advogados particulares:" in counsel
    assert "Dr. Advogado Um - OAB/RN 111 (em defesa de Réu Alfa)" in counsel
    assert "Dr. Advogado Dois - OAB/RN 222 (em defesa de Réu Beta)" in counsel


