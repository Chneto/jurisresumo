"""Unit tests for UC-2: Extended Procedural Hierarchy CPP (Júri, Custódia, Protetivas, ANPP, AIJ, PAnP).

Desenvolvido por FChNeto.
"""

import pytest
import docx

from app.core.models import (
    Defendant,
    DocumentCategory,
    HearingSummaryData,
    PJeDocument,
)
from app.core.jev_decision_engine import (
    ClassificationHierarchyCPP,
    JEVDecisionEngine,
)
from app.engines.offline_engine import OfflineExtractionEngine
from app.generators.docx_generator import generate_docx_summary


class TestExtendedHierarchyClassification:
    """Tests the precedence hierarchy of ClassificationHierarchyCPP and JEVDecisionEngine."""

    def test_priority_1_juri_pronuncia_and_plenario(self):
        """Pauta do júri e pronúncia devem classificar como JÚRI."""
        engine = JEVDecisionEngine.get_instance()

        # Contexto com pronúncia e sessão plenária
        h_comb = "despacho designando sessão plenária do tribunal do júri com sorteio dos jurados"
        catalog = [
            PJeDocument(
                doc_id="101",
                date_str="01/02/2026",
                doc_name="Decisão de Pronúncia",
                doc_type="Decisão",
                start_page=1,
                end_page=5,
            )
        ]

        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="PROCESSO_0800123_SESSAO_JURI.PDF",
            pje_catalog=catalog,
        )
        assert act == "JÚRI"

        # Via classify_text_block
        res = engine.classify_text_block(
            text="Fica designada a Sessão Plenária do Tribunal do Júri para o dia 15 de outubro de 2026.",
            doc_title="Despacho pauta de júri",
        )
        assert res.category == DocumentCategory.SESSAO_JURI

    def test_priority_2_custodia_apf_sem_denuncia(self):
        """Auto de Prisão em Flagrante sem oferecimento/recebimento de denúncia deve ser CUSTÓDIA."""
        engine = JEVDecisionEngine.get_instance()

        h_comb = "termo de audiência de custódia apresentação do preso em flagrante delito"
        catalog = [
            PJeDocument(
                doc_id="201",
                date_str="05/03/2026",
                doc_name="Auto de Prisão em Flagrante - APF 123/2026",
                doc_type="Auto de Prisão em Flagrante",
                start_page=1,
                end_page=10,
            )
        ]

        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="0802345_CUSTODIA.PDF",
            pje_catalog=catalog,
            has_received_denuncia=False,
            has_denuncia=False,
        )
        assert act == "CUSTÓDIA"

        # Via classify_text_block
        res = engine.classify_text_block(
            text="Designo audiência de custódia para apresentação da pessoa presa em flagrante.",
            doc_title="Decisão de Custódia",
        )
        assert res.category == DocumentCategory.AUDIENCIA_CUSTODIA

    def test_priority_2_custodia_blocked_if_denuncia_received(self):
        """Se houver denúncia recebida com instrução, NÃO deve classificar como Custódia."""
        h_comb = "auto de prisão em flagrante e designação de audiência de instrução e julgamento"
        catalog = [
            PJeDocument(
                doc_id="201",
                date_str="05/03/2026",
                doc_name="Denúncia",
                doc_type="Denúncia",
                start_page=1,
                end_page=4,
            ),
            PJeDocument(
                doc_id="202",
                date_str="10/03/2026",
                doc_name="Decisão Recebimento Denúncia",
                doc_type="Decisão",
                start_page=5,
                end_page=6,
            ),
        ]

        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="PROCESSO_0802345.PDF",
            pje_catalog=catalog,
            has_received_denuncia=True,
            has_denuncia=True,
            has_aij_mention=True,
        )
        assert act == "AIJ"

    def test_priority_3_revocation_anpp_forces_aij(self):
        """Revogação de ANPP ou descumprimento força obrigatoriamente AIJ."""
        h_comb = "decisão revogando o acordo de não persecução penal por descumprimento das condições"
        anpp_rev = PJeDocument(
            doc_id="301",
            date_str="15/04/2026",
            doc_name="Decisão Revogação ANPP",
            doc_type="Decisão",
            start_page=1,
            end_page=3,
        )

        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="PROCESSO_ANPP_REVOGADO.PDF",
            anpp_rev_doc=anpp_rev,
        )
        assert act == "AIJ"

    def test_priority_4_panp_art_366_cpp(self):
        """Produção Antecipada de Provas nos termos do art. 366 do CPP."""
        engine = JEVDecisionEngine.get_instance()

        h_comb = "decisão determinando a suspensão do processo e produção antecipada de provas orais art 366 cpp"
        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="0803456_PANP.PDF",
        )
        assert act == "PAnP"

        res = engine.classify_text_block(
            text="Determino a suspensão do feito e a produção antecipada de provas com fulcro no art. 366 do CPP.",
            doc_title="Decisão suspensão art 366",
        )
        assert res.category == DocumentCategory.DECISAO_PANP

    def test_priority_5_anpp_homologatorio(self):
        """Audiência designada exclusivamente para homologação de ANPP."""
        engine = JEVDecisionEngine.get_instance()

        h_comb = "designo audiência para homologação de acordo de não persecução penal nos termos do art. 28-A do CPP"
        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="0804567_ANPP.PDF",
            has_anpp_rejection=False,
        )
        assert act == "ANPP"

        res = engine.classify_text_block(
            text="Audiência para fins de homologação do Acordo de Não Persecução Penal - art. 28-A do CPP.",
            doc_title="Termo ANPP",
        )
        assert res.category == DocumentCategory.DECISAO_ANPP

    def test_priority_6_aij_ordinaria(self):
        """Audiência de Instrução e Julgamento ordinária."""
        engine = JEVDecisionEngine.get_instance()

        h_comb = "designo audiência de instrução e julgamento para oitiva das testemunhas e interrogatório do réu"
        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="0805678_AIJ.PDF",
            has_received_denuncia=True,
            has_aij_mention=True,
        )
        assert act == "AIJ"

        res = engine.classify_text_block(
            text="Recebo a denúncia e designo audiência de instrução e julgamento.",
            doc_title="Decisão AIJ",
        )
        assert res.category == DocumentCategory.DECISAO_AIJ

    def test_priority_7_protetiva_maria_da_penha_justificacao(self):
        """Audiência de justificação da Lei Maria da Penha (art. 16 da Lei 11.340/06)."""
        engine = JEVDecisionEngine.get_instance()

        h_comb = "audiência de justificação para retratação da vítima nos termos do art. 16 da lei 11.340 medida protetiva"
        act = ClassificationHierarchyCPP.classify(
            h_comb=h_comb,
            fn_upper="PROCESSO_0806789_PROTETIVA.PDF",
        )
        assert act == "JUSTIFICAÇÃO / PROTETIVA"

        res = engine.classify_text_block(
            text="Designo audiência de justificação prevista no art. 16 da Lei 11.340/2006 (Lei Maria da Penha).",
            doc_title="Despacho Lei Maria da Penha",
        )
        assert res.category == DocumentCategory.AUDIENCIA_PROTETIVA


class TestExtendedHierarchyDocxCallouts:
    """Tests the rendering of specialized yellow callouts in Word DOCX summaries."""

    def test_docx_callout_juri(self):
        """DOCX para Júri deve renderizar callout 'Sessão Plenária do Tribunal do Júri'."""
        data = HearingSummaryData(
            case_number="0801111-22.2026.8.20.5001",
            act_type="JÚRI",
            hearing_datetime="18.11.26 às 08h30min",
            prosecutor="Dr. Promotor Titular do Júri",
            defendants=[Defendant(name="João Réu Pronunciado", status="Preso preventivo")],
            defense_counsel="Dr. Advogado do Júri - OAB/RN 1234",
            closure_text="Cordial e respeitosamente,",
        )

        docx_bytes = generate_docx_summary(data)
        assert len(docx_bytes) > 0

        # Carregar DOCX e verificar parágrafos
        import io
        doc = docx.Document(io.BytesIO(docx_bytes))
        texts = [p.text for p in doc.paragraphs]

        # Callout esperado: "08h30min - 0801111-22.2026.8.20.5001 - Sessão Plenária do Tribunal do Júri"
        callout_matches = [t for t in texts if "Sessão Plenária do Tribunal do Júri" in t]
        assert len(callout_matches) == 1
        assert "08h30min" in callout_matches[0]
        assert "0801111-22.2026.8.20.5001" in callout_matches[0]

    def test_docx_callout_custodia(self):
        """DOCX para Custódia deve renderizar callout 'Audiência de Custódia'."""
        data = HearingSummaryData(
            case_number="0802222-33.2026.8.20.5001",
            act_type="CUSTÓDIA",
            hearing_datetime="10.10.26 às 14h00min",
            prosecutor="Dr. Promotor Plantonista",
            defendants=[Defendant(name="Flagranteado Silva", status="Preso em flagrante")],
            defense_counsel="Assistido pela Defensoria Pública",
            closure_text="Cordial e respeitosamente,",
        )

        docx_bytes = generate_docx_summary(data)
        import io
        doc = docx.Document(io.BytesIO(docx_bytes))
        texts = [p.text for p in doc.paragraphs]

        callout_matches = [t for t in texts if "Audiência de Custódia" in t]
        assert len(callout_matches) == 1
        assert "14h00min" in callout_matches[0]

    def test_docx_callout_protetiva(self):
        """DOCX para Protetivas deve renderizar callout 'Audiência de Justificação (Lei Maria da Penha)'."""
        data = HearingSummaryData(
            case_number="0803333-44.2026.8.20.5001",
            act_type="JUSTIFICAÇÃO / PROTETIVA",
            hearing_datetime="22.10.26 às 11h15min",
            prosecutor="Dr. Promotor da Mulher",
            defendants=[Defendant(name="Autor do Fato", status="Em liberdade com cautelares")],
            defense_counsel="Assistido pela Defensoria Pública",
            closure_text="Cordial e respeitosamente,",
        )

        docx_bytes = generate_docx_summary(data)
        import io
        doc = docx.Document(io.BytesIO(docx_bytes))
        texts = [p.text for p in doc.paragraphs]

        callout_matches = [t for t in texts if "Audiência de Justificação (Lei Maria da Penha)" in t]
        assert len(callout_matches) == 1
        assert "11h15min" in callout_matches[0]


class TestOfflineEngineDelegationExtendedHierarchy:
    """Tests that OfflineExtractionEngine._determine_act_type delegates smoothly."""

    def test_offline_engine_classifies_extended_types(self):
        engine = OfflineExtractionEngine()

        # Mock Júri via pje_catalog
        juri_catalog = [
            PJeDocument(
                doc_id="888",
                date_str="01/01/2026",
                doc_name="Decisão Pronúncia",
                doc_type="Decisão",
                start_page=1,
                end_page=5,
            )
        ]
        act_juri = engine._determine_act_type(
            filename="PROCESSO_080001.PDF",
            hearing_text="Pauta do Júri - Sessão de Julgamento",
            hearing_doc_name="Edital de Convocação do Júri",
            doc=None,
            denuncia_doc=None,
            pje_catalog=juri_catalog,
        )
        assert act_juri == "JÚRI"

        # Mock Protetiva
        act_prot = engine._determine_act_type(
            filename="PROCESSO_080002.PDF",
            hearing_text="Audiência de Justificação nos termos do artigo 16 da Lei Maria da Penha",
            hearing_doc_name="Despacho de Designação",
            doc=None,
            denuncia_doc=None,
            pje_catalog=[],
        )
        assert act_prot == "JUSTIFICAÇÃO / PROTETIVA"
